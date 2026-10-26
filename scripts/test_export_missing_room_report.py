import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from export_missing_room_report import export, group_exceptions, match_rules
from validate_excel import save_excel


class RoomExceptionTests(unittest.TestCase):
    def test_prefix_uses_selected_field_and_case(self):
        rules = [{"id": "asp", "field": "Type", "prefix": "Aspirasjonsrør"}]
        self.assertEqual(match_rules({"Familienavn": "Conduit with Fittings", "Type": "ASPIRASJONSRØR 25"}, rules), ["asp"])
        self.assertEqual(match_rules({"Familienavn": "Conduit with Fittings", "Type": "Vanlig rør"}, rules), [])

    def test_group_once_per_family_and_type(self):
        members = [dict(Familienavn="Familie", Type="Type", ElementId=str(identifier), UniqueId=f"uid-{identifier}")
                   for identifier in (1, 2)]
        grouped = group_exceptions(members)
        self.assertEqual(len(grouped), 1)
        self.assertEqual(grouped[0][2], 2)
        self.assertEqual(grouped[0][5], "1; 2")
        with self.assertRaises(ValueError):
            group_exceptions(members + members)

    def test_export_preserves_source_and_supports_grouped_input(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, output, repeated = (root / name for name in ("source.xlsx", "output.xlsx", "repeat.xlsx"))
            config = root / "rules.json"
            config.write_text(json.dumps(dict(schemaVersion=1, modelPath="BAS", enterprise="K5B",
                                             parameters=["Romnavn", "dRofus_RoomfunctionNo"],
                                             rules=[dict(id="type", field="Type", prefix="Asp")])), encoding="utf-8")
            book = Workbook()
            main = book.active
            main.title = "Manglende romdata"
            headers = ["ElementId", "UniqueId", "Familienavn", "Type", "FOB_Entreprise", "Avvik"]
            main.append(headers)
            main.append(["1", "uid-1", "Rør", "Asp", "K5B", "Mangler romnavn"])
            main.append(["2", "uid-2", "Rør", "Vanlig", "K5B", "Mangler romnavn"])
            main.append(["3", "uid-3", "Rør", "Asp", "K5B", "Mangler romnavn"])
            main["A3"].comment = Comment("Behold notat", "Bruker")
            book.create_sheet("Unntak").append(headers)
            summary = book.create_sheet("Oppsummering")
            summary.append(["Kontroll", "Verdi"])
            summary.append(["Modellsti", "BAS"])
            save_excel(book, source)
            book.close()
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            self.assertEqual(export(source, output, config), (1, 2, 1))
            self.assertEqual(export(output, repeated, config), (1, 0, 1))
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)
            check = load_workbook(output)
            self.assertEqual(check["Unntak"].max_row, 2)
            self.assertEqual(check["Manglende romdata"]["A2"].comment.text, "Behold notat")
            check.close()
            with self.assertRaises(FileExistsError):
                export(source, output, config)


if __name__ == "__main__":
    unittest.main()