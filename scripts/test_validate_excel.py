import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn

from export_confirmed_routes import export
from validate_excel import save_excel, validate_excel


class ExcelTableValidationTests(unittest.TestCase):
    def make_workbook(self, path, *, count=2, ids=(1, 2), names=("Id", "SpaceNavn"), filter_ref="A1:B2"):
        main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with ZipFile(path, "w") as archive:
            archive.writestr(
                "xl/worksheets/sheet1.xml",
                f'<worksheet xmlns="{main}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Id</t></is></c>'
                '<c r="B1" t="inlineStr"><is><t>SpaceNavn</t></is></c></row></sheetData>'
                '<tableParts count="1"><tablePart r:id="rId1"/></tableParts></worksheet>',
            )
            archive.writestr(
                "xl/worksheets/_rels/sheet1.xml.rels",
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Target="../tables/table1.xml"/></Relationships>',
            )
            columns = "".join(f'<tableColumn id="{i}" name="{n}"/>' for i, n in zip(ids, names))
            archive.writestr(
                "xl/tables/table1.xml",
                f'<table xmlns="{main}" displayName="TestTable" ref="A1:B2">'
                f'<autoFilter ref="{filter_ref}"/><tableColumns count="{count}">{columns}</tableColumns></table>',
            )

    def test_valid_table(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            self.make_workbook(path)
            self.assertEqual(validate_excel(path), 1)

    def test_invalid_table_metadata(self):
        cases = [
            {"count": 3},
            {"ids": (1, 1)},
            {"ids": (0, 2)},
            {"names": ("Id", "Id")},
            {"names": ("Id", "WrongHeader")},
            {"filter_ref": "A1:A2"},
            {"count": 4, "ids": (1, 2, 1, 2), "names": ("Id", "SpaceNavn", "Column1", "Column2")},
        ]
        with tempfile.TemporaryDirectory() as directory:
            for case in cases:
                with self.subTest(case=case):
                    path = Path(directory) / "test.xlsx"
                    self.make_workbook(path, **case)
                    with self.assertRaises(ValueError):
                        validate_excel(path)

    def test_inline_string_content(self):
        main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        cases = [
            ('<c r="A1" t="inlineStr"/>', False),
            ('<c r="A1" t="inlineStr"><is><t/></is></c>', True),
            ('<c r="A1" s="1"/>', True),
            ('<c r="A1" t="inlineStr"><is><t>Space</t></is></c>', True),
        ]
        with tempfile.TemporaryDirectory() as directory:
            for cell, valid in cases:
                with self.subTest(cell=cell):
                    path = Path(directory) / "test.xlsx"
                    with ZipFile(path, "w") as archive:
                        archive.writestr(
                            "xl/worksheets/sheet1.xml",
                            f'<worksheet xmlns="{main}"><sheetData><row r="1">{cell}</row></sheetData></worksheet>',
                        )
                    if valid:
                        self.assertEqual(validate_excel(path), 0)
                    else:
                        with self.assertRaisesRegex(ValueError, "inlineStr cell missing is content"):
                            validate_excel(path)


class ExcelSaveTests(unittest.TestCase):
    def workbook(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["ElementId", "Romnavn", "Merknad"])
        sheet.append(["001", "", "=tekst, ikke formel"])
        sheet["C2"].data_type = "s"
        sheet.append(["002", None, "æøå"])
        sheet["B3"].data_type = "s"
        sheet.append(["003", "Rom", "=1+1"])
        sheet.freeze_panes = "A2"
        return book

    def test_save_normalizes_blanks_and_restores_view(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            book = self.workbook()
            book.views = []
            book.save(path)
            with self.assertRaisesRegex(ValueError, "missing workbookView"):
                validate_excel(path)
            save_excel(book, path)
            self.assertEqual(validate_excel(path), 0)
            check = load_workbook(path)
            self.assertEqual(len(check.views), 1)
            self.assertEqual(check.active.freeze_panes, "A2")
            self.assertIsNone(check.active["B2"].value)
            self.assertIsNone(check.active["B3"].value)
            self.assertEqual(check.active["A2"].value, "001")
            self.assertEqual(check.active["C2"].value, "=tekst, ikke formel")
            self.assertEqual(check.active["C2"].data_type, "s")
            self.assertEqual(check.active["C3"].value, "æøå")
            self.assertEqual(check.active["C4"].data_type, "f")
            check.close()
            book.close()

    def test_save_removes_only_duplicate_sheet_filter(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            book = self.workbook()
            for reference in ("B2", "B3"):
                book.active[reference].value = None
                book.active[reference].data_type = "n"
            book.active.add_table(Table(displayName="Romdata", ref="A1:C4"))
            book.active.auto_filter.ref = "A1:C4"
            book.save(path)
            with self.assertRaisesRegex(ValueError, "worksheet autoFilter overlaps"):
                validate_excel(path)
            save_excel(book, path)
            self.assertEqual(validate_excel(path), 1)
            self.assertIsNone(book.active.auto_filter.ref)
            self.assertEqual(book.active.tables["Romdata"].autoFilter.ref, "A1:C4")
            book.close()

    def test_save_rejects_partial_filter_overlap(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            book = self.workbook()
            book.active.add_table(Table(displayName="Romdata", ref="A1:C4"))
            book.active.auto_filter.ref = "A1:B4"
            with self.assertRaisesRegex(ValueError, "partially overlaps"):
                save_excel(book, path)
            self.assertFalse(path.exists())
            book.close()

    def test_save_preserves_active_filter_by_rejecting_overlap(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            book = self.workbook()
            book.active.add_table(Table(displayName="Romdata", ref="A1:C4"))
            book.active.auto_filter.ref = "A1:C4"
            book.active.auto_filter.add_filter_column(0, ["001"])
            with self.assertRaisesRegex(ValueError, "active worksheet autoFilter"):
                save_excel(book, path)
            self.assertEqual(book.active.auto_filter.filterColumn[0].filters.filter, ["001"])
            self.assertFalse(path.exists())
            book.close()

    def test_table_filter_excludes_total_row(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            book = self.workbook()
            table = Table(displayName="Romdata", ref="A1:C4", totalsRowCount=1,
                          autoFilter=AutoFilter(ref="A1:C3"), tableColumns=[
                              TableColumn(id=index, name=cell.value)
                              for index, cell in enumerate(book.active[1], 1)
                          ])
            book.active.add_table(table)
            save_excel(book, path)
            self.assertEqual(validate_excel(path), 1)
            book.close()

    def test_invalid_table_does_not_replace_existing_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            original = b"Existing report must survive failed validation"
            path.write_bytes(original)
            book = Workbook()
            book.active.append(["Id", "Id"])
            book.active.append(["1", "2"])
            book.active.add_table(Table(displayName="InvalidHeaders", ref="A1:B2"))
            with self.assertRaisesRegex(ValueError, "duplicate or empty column names"):
                save_excel(book, path)
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(Path(directory).iterdir()), [path])
            book.close()


class ConfirmedRouteExportTests(unittest.TestCase):
    def test_export_preserves_manual_formula_and_literal_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "config").mkdir()
            (root / "logs" / "reports").mkdir(parents=True)
            registry = [dict(Route="R1", RepresentativeUniqueId="uid-1", UniqueId=f"uid-{index}",
                             ElementId=str(index), ExpectedOccurrenceId=str(100 + index),
                             WidthMm="300", HeightMm="60", ModelPath="BAS", ViewUniqueId="view-1")
                        for index in (1, 2)]
            snapshot = [dict(UniqueId=f"uid-{index}", ElementId=str(index), OccurrenceId=str(100 + index),
                             Hidden="False" if index == 1 else "True", Mark="MERKE-001", Sequence="001",
                             RoomId="1", RoomNumber="E1S.01", RoomName="Rom", TypeName="Kabelstige",
                             C4="False", ItemId="1") for index in (1, 2)]
            for path, rows in ((root / "config" / "DrofusConfirmedRoutes.tsv", registry),
                               (root / "snapshot.tsv", snapshot)):
                with path.open("w", encoding="utf-8", newline="") as stream:
                    writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter="\t")
                    writer.writeheader()
                    writer.writerows(rows)
            output = root / "report.xlsx"
            with patch("export_confirmed_routes.ROOT", root):
                export(root / "snapshot.tsv", output)
                book = load_workbook(output)
                book["Segmenter"]["Z2"] = "=1+1"
                book["Segmenter"]["Z3"] = "=manuell tekst"
                book["Segmenter"]["Z3"].data_type = "s"
                save_excel(book, output)
                book.close()
                export(root / "snapshot.tsv", output)
            check = load_workbook(output)
            self.assertEqual(check["Segmenter"]["Z2"].value, "=1+1")
            self.assertEqual(check["Segmenter"]["Z2"].data_type, "f")
            self.assertEqual(check["Segmenter"]["Z3"].value, "=manuell tekst")
            self.assertEqual(check["Segmenter"]["Z3"].data_type, "s")
            self.assertEqual(len(list((root / "logs" / "reports").glob("*.xlsx"))), 1)
            self.assertEqual(validate_excel(output), 3)
            check.close()


if __name__ == "__main__":
    unittest.main()
