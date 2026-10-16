"""Export registered route links and a fresh Revit snapshot to a cumulative workbook."""

import argparse
import csv
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
HEADERS = [
    "Strekk", "Rolle", "RepresentantId", "ElementId", "UniqueId",
    "RepresentativeUniqueId", "Typenavn", "C4", "Bredde mm", "Høyde mm",
    "Forventet occurrence", "Faktisk occurrence", "ItemID", "Skjult",
    "Merkestreng", "Sekvensnummer", "RomID", "Romnummer", "Romnavn",
    "Representantmerke", "Representantsekvens", "Representantromnummer",
    "Status", "Modellsti", "ViewUniqueId", "Merknad",
]


def read_tsv(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def valid(value):
    return bool(value.strip()) and value != "--" and not value.startswith("Error:")


def add_sheet(workbook, name, headers, rows, table_name):
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([None if value == "" else value for value in row])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            55, max(14, max(len(str(cell.value or "")) for cell in column) + 2)
        )


def export(snapshot, output):
    registry = read_tsv(ROOT / "config" / "DrofusConfirmedRoutes.tsv")
    fresh = read_tsv(snapshot)
    by_uid = {row["UniqueId"]: row for row in fresh}
    if not registry or len(by_uid) != len(fresh):
        raise ValueError("Empty registry or duplicate snapshot UniqueId")
    for field in ("UniqueId", "ElementId", "ExpectedOccurrenceId"):
        if len({row[field] for row in registry}) != len(registry):
            raise ValueError(f"Duplicate registered {field}")
    if set(by_uid) != {row["UniqueId"] for row in registry}:
        raise ValueError("Snapshot must contain exactly all registered segments")
    groups = defaultdict(list)
    for row in registry:
        groups[row["Route"]].append(row)
    notes = {}
    if output.exists():
        previous = load_workbook(output)
        try:
            for values in list(previous["Segmenter"].values)[1:]:
                notes[values[4]] = values[25]
            if set(notes) - set(by_uid):
                raise ValueError("Previously reported segments are missing; refusing to drop notes")
        finally:
            previous.close()
    segments = []
    summaries = []
    for route, members in groups.items():
        owners = {row["RepresentativeUniqueId"] for row in members}
        if len(owners) != 1 or len(members) < 2:
            raise ValueError(f"Invalid route membership: {route}")
        owner_uid = next(iter(owners))
        if sum(row["UniqueId"] == owner_uid for row in members) != 1:
            raise ValueError(f"Missing registered representative: {route}")
        owner = by_uid[owner_uid]
        statuses = []
        for member in members:
            current = by_uid[member["UniqueId"]]
            if member["ElementId"] != current["ElementId"]:
                raise ValueError("ElementId/UniqueId mismatch")
            issues = []
            representative = member["UniqueId"] == owner_uid
            if current["OccurrenceId"] != member["ExpectedOccurrenceId"]:
                issues.append("Occurrence-ID endret – avklar")
            if current["Hidden"] != ("False" if representative else "True"):
                issues.append("Synlighet avviker")
            if not valid(current["Mark"]) or not current["Sequence"].isascii() or not current["Sequence"].isdigit() or int(current["Sequence"]) <= 0:
                issues.append("Mangler gyldig merke/sekvens")
            if not valid(owner["Mark"]) or not owner["Sequence"].isascii() or not owner["Sequence"].isdigit() or int(owner["Sequence"]) <= 0:
                issues.append("Representant mangler gyldig merkekilde")
            elif current["Mark"] != owner["Mark"] or current["Sequence"] != owner["Sequence"]:
                issues.append("Merke/sekvens avviker fra representant")
            if int(current["RoomId"]) <= 0 or not valid(current["RoomNumber"]):
                issues.append("Mangler rominfo")
            elif current["RoomId"] != owner["RoomId"] or current["RoomNumber"] != owner["RoomNumber"]:
                issues.append("Rominfo avviker – avklar")
            status = "; ".join(issues) or "Kontrollerte felt samsvarer i snapshot"
            statuses.append(status)
            segments.append([
                route, "Representant" if representative else "Segment",
                owner["ElementId"], current["ElementId"], member["UniqueId"], owner_uid,
                current["TypeName"], current["C4"], float(member["WidthMm"]), float(member["HeightMm"]),
                int(member["ExpectedOccurrenceId"]), int(current["OccurrenceId"]), int(current["ItemId"]),
                current["Hidden"], current["Mark"], current["Sequence"], int(current["RoomId"]),
                current["RoomNumber"], current["RoomName"], owner["Mark"], owner["Sequence"],
                owner["RoomNumber"], status, member["ModelPath"], member["ViewUniqueId"], notes.get(member["UniqueId"]),
            ])
        summaries.append([route, owner["ElementId"], owner_uid, len(members),
                          owner["OccurrenceId"], owner["Mark"], owner["Sequence"],
                          owner["RoomNumber"], " | ".join(dict.fromkeys(statuses))])
    book = Workbook()
    book.remove(book.active)
    add_sheet(book, "Les meg", ["Tema", "Beskrivelse"], [
        ["Grunnlag", "Brukerbekreftede koblinger i config\\DrofusConfirmedRoutes.tsv"],
        ["Snapshot", str(snapshot.resolve())],
        ["Generert", datetime.now().isoformat(timespec="seconds")],
        ["Autoritet", "Excel er rapport; TSV er medlemsregister. Merkestreng/rominfo er ferske snapshotverdier, ikke permanent fasit."],
        ["Utfyllingskrav", "Alle segmenter skal ha merke og sekvens. Skjuling opphever ikke avvik."],
        ["Etter synk", "Hent registrerte UniqueId-er og les ferske verdier fra representanten. Ikke kopier occurrence-ID."],
        ["Merknader", "Skriv manuelle notater i Merknad på Segmenter; eksport bevarer dem etter UniqueId."],
        ["Kjøring", "Ingen automatisk synk eller modellendring. Rapporten må fornyes fra et ferskt modelluttrekk."],
    ], "ReadMe")
    add_sheet(book, "Strekk", ["Strekk", "RepresentantId", "RepresentativeUniqueId",
                              "Antall segmenter", "Occurrence", "Merkestreng", "Sekvensnummer",
                              "Romnummer", "Status"], summaries, "Routes")
    add_sheet(book, "Segmenter", HEADERS, segments, "Segments")
    add_sheet(book, "Avvik", HEADERS, [row for row in segments if row[22] != "Kontrollerte felt samsvarer i snapshot"], "Exceptions")
    if output.exists():
        backup = ROOT / "logs" / "reports" / f"BAS bekreftede strekk før oppdatering {datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
        shutil.copy2(output, backup)
    book.save(output)
    book.close()
    check = load_workbook(output)
    try:
        assert check["Segmenter"].max_row - 1 == len(registry)
        assert {row[4] for row in list(check["Segmenter"].values)[1:]} == set(by_uid)
        assert check["Strekk"].max_row - 1 == len(groups)
    finally:
        check.close()
    print(f"{output}: {len(groups)} strekk, {len(segments)} segmenter")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("snapshot", type=Path)
    parser.add_argument("--output", type=Path, default=ROOT / "config" / "DrofusConfirmedRoutes.xlsx")
    arguments = parser.parse_args()
    export(arguments.snapshot, arguments.output)
