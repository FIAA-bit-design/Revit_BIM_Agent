# -*- coding: utf-8 -*-
"""Apply confirmed room-report exceptions without changing the source workbook."""

import argparse
from collections import Counter
from copy import copy
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from validate_excel import save_excel


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RULES = ROOT / "config" / "RoomInformationExceptions.json"


def load_rules(path):
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("schemaVersion") != 1 or not data.get("modelPath") or data.get("enterprise") != "K5B":
        raise ValueError("Invalid room-exception scope")
    if data.get("parameters") != ["Romnavn", "dRofus_RoomfunctionNo"]:
        raise ValueError("Exceptions apply only to room name/number reporting")
    identities = set()
    for rule in data["rules"]:
        if rule["id"] in identities or rule["field"] not in ("Familienavn", "Type") or not rule["prefix"].strip():
            raise ValueError("Duplicate, empty or invalid exception rule")
        identities.add(rule["id"])
    return data


def match_rules(row, rules):
    return [rule["id"] for rule in rules
            if str(row.get(rule["field"]) or "").casefold().startswith(rule["prefix"].casefold())]


def records(sheet):
    values = list(sheet.values)
    headers = list(values[0])
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate worksheet headers")
    return headers, [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]


def previous_exceptions(sheet):
    headers, rows = records(sheet)
    if "ElementId" in headers:
        return [dict(row, Unntaksgrunnlag=row.get("Avvik") or "Tidligere bekreftet unntak") for row in rows]
    members = []
    for row in rows:
        identifiers = str(row["ElementId-er"]).split("; ")
        unique_ids = str(row["UniqueId-er"]).split("; ")
        if len(identifiers) != len(unique_ids) or len(identifiers) != int(row["Antall elementer"]):
            raise ValueError("Invalid grouped exception membership")
        for identifier, unique_id in zip(identifiers, unique_ids):
            members.append(dict(row, ElementId=identifier, UniqueId=unique_id))
    return members


def group_exceptions(rows):
    groups = {}
    seen = set()
    for row in rows:
        if not row.get("UniqueId") or row["UniqueId"] in seen:
            raise ValueError("Missing or duplicate exception UniqueId")
        seen.add(row["UniqueId"])
        key = (str(row.get("Familienavn") or "").casefold(), str(row.get("Type") or "").casefold())
        groups.setdefault(key, []).append(row)
    result = []
    for key in sorted(groups):
        members = sorted(groups[key], key=lambda row: int(row["ElementId"]))
        result.append([
            members[0].get("Familienavn"), members[0].get("Type"), len(members),
            "; ".join(sorted({str(row.get("Revit-kategori") or "") for row in members})),
            "; ".join(sorted({str(row.get("Unntaksgrunnlag") or "Tidligere bekreftet unntak") for row in members})),
            "; ".join(str(row["ElementId"]) for row in members),
            "; ".join(row["UniqueId"] for row in members)
        ])
    if any(isinstance(value, str) and len(value) > 32767 for row in result for value in row):
        raise ValueError("Exception membership exceeds Excel cell capacity; refusing truncation")
    return result


def style_sheet(sheet, table_name):
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24536B")
    for index, cell in enumerate(sheet[1], 1):
        sheet.column_dimensions[get_column_letter(index)].width = 18 if cell.value in ("ElementId", "Antall elementer") else 40
    if sheet.max_row > 1:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def add_rows(sheet, headers, rows):
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str) and cell.value:
                cell.data_type = "s"


def export(source, output, rules_path=DEFAULT_RULES):
    if output.exists() or source.resolve() == output.resolve():
        raise FileExistsError("A new output file is required; source and previous reports are preserved")
    config = load_rules(rules_path)
    original = load_workbook(source)
    result = Workbook()
    try:
        metadata = dict(original["Oppsummering"].iter_rows(min_row=2, max_col=2, values_only=True))
        if metadata.get("Modellsti") != config["modelPath"]:
            raise ValueError("Source model does not match exception scope")
        headers, rows = records(original["Manglende romdata"])
        inherited = previous_exceptions(original["Unntak"])
        if any(str(row.get("FOB_Entreprise", config["enterprise"]) or "").strip().casefold() != config["enterprise"].casefold()
               for row in rows + inherited):
            raise ValueError("Source contains another enterprise")
        if len({row["UniqueId"] for row in rows + inherited}) != len(rows) + len(inherited):
            raise ValueError("Duplicate source membership")
        kept, new_exceptions = [], []
        for row in rows:
            matches = match_rules(row, config["rules"])
            if matches:
                new_exceptions.append(dict(row, Unntaksgrunnlag="; ".join(matches)))
            else:
                kept.append(row)
        main = result.active
        main.title = "Manglende romdata"
        main.append(headers)
        source_rows = {original["Manglende romdata"].cell(index, headers.index("UniqueId") + 1).value: index
                       for index in range(2, original["Manglende romdata"].max_row + 1)}
        for row in kept:
            source_cells = original["Manglende romdata"][source_rows[row["UniqueId"]]]
            main.append([row[header] for header in headers])
            for target_cell, source_cell in zip(main[main.max_row], source_cells):
                target_cell.data_type = source_cell.data_type
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.number_format = source_cell.number_format
                target_cell.comment = copy(source_cell.comment)
                target_cell.hyperlink = copy(source_cell.hyperlink)
        style_sheet(main, "FilteredRoomDeviations")
        grouped = group_exceptions(inherited + new_exceptions)
        exceptions = result.create_sheet("Unntak")
        add_rows(exceptions, ["Familienavn", "Type", "Antall elementer", "Revit-kategori", "Unntaksgrunnlag", "ElementId-er", "UniqueId-er"], grouped)
        style_sheet(exceptions, "GroupedRoomExceptions")
        filters = result.create_sheet("Filterregler")
        add_rows(filters, ["Regel", "Felt", "Prefiks", "Gule kildeceller"],
                 [[rule["id"], rule["field"], rule["prefix"], "; ".join(rule.get("sourceCells", []))] for rule in config["rules"]])
        style_sheet(filters, "RoomExceptionRules")
        summary = result.create_sheet("Oppsummering")
        counts = Counter(row["Avvik"] for row in kept)
        add_rows(summary, ["Kontroll", "Resultat"], [
            ["Modellsti", config["modelPath"]], ["Entreprise", config["enterprise"]],
            ["Datagrunnlag", "Filtrert fra eksisterende rapport; ingen ny modellkontroll eller modellendring."],
            ["Kilderapport", str(source.resolve())], ["Unntaksregler", str(rules_path.resolve())],
            ["Tidspunkt fra Revit-host", metadata.get("Tidspunkt fra Revit-host")],
            ["K5B-elementer", metadata.get("K5B-elementer")],
            ["Avvik før ny filtrering", len(rows)], ["Nye unntakselementer", len(new_exceptions)],
            ["Tidligere unntakselementer", len(inherited)], ["Unntakselementer totalt", len(inherited) + len(new_exceptions)],
            ["Unike familie-/typekombinasjoner i unntak", len(grouped)], ["Avvik", len(kept)],
            ["Mangler begge felt", counts["Mangler romnavn og romnummer"]],
            ["Mangler bare romnavn", counts["Mangler romnavn"]],
            ["Mangler bare romnummer", counts["Mangler romnummer"]]
        ])
        style_sheet(summary, "RoomReportSummary")
        summary.column_dimensions["B"].width = 110
        save_excel(result, output)
        check = load_workbook(output, read_only=True)
        try:
            actual = list(check["Manglende romdata"].iter_rows(min_row=2, values_only=True))
            expected = [tuple(row[header] if row[header] != "" else None for header in headers) for row in kept]
            expected_groups = [tuple(value if value != "" else None for value in row) for row in grouped]
            if actual != expected or list(check["Unntak"].values)[1:] != expected_groups:
                raise ValueError("Saved report differs from source data or grouped exceptions")
        finally:
            check.close()
        print(f"{output}: {len(kept)} avvik; {len(new_exceptions)} nye unntak; "
              f"{len(inherited) + len(new_exceptions)} unntakselementer i {len(grouped)} familie-/typekombinasjoner")
        return len(kept), len(new_exceptions), len(grouped)
    finally:
        original.close()
        result.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--rules", type=Path, default=DEFAULT_RULES)
    arguments = parser.parse_args()
    export(arguments.source, arguments.output, arguments.rules)